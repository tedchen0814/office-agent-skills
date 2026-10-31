"""template_deck_lib.py：用公司 .pptx 範本（母片）做大量文字頁的排版輔助。

提供：自動字級（優先讓每條文字一行寫完）、條列／表格／程式碼框／提示框、欄位頁、
空白處自動補圖（auto_img）、表格儲存格的點擊跳轉（TABLES 記錄座標）。
使用前先改下面 CONFIG：範本路徑、版型索引、配色。用法見 references/template-deck.md。
"""
import os
import glob, math, re
from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ── CONFIG（依範本修改）──
ORG, GRY, DK, LT, WH, BD = "D64A00", "7F8286", "333333", "F2F2F2", "FFFFFF", "D5D5D5"  # 強調色、灰、深灰、淺灰、白、框線
FONT, MONO = "Microsoft JhengHei", "Consolas"
rgb = lambda h: RGBColor.from_string(h)
TOK = re.compile(r"(\*\*.+?\*\*|`.+?`)")

TEMPLATE = os.environ.get("DECK_TEMPLATE", "template.pptx")
prs = Presentation(TEMPLATE)
layouts = [l for m in prs.slide_masters for l in m.slide_layouts]
L_COVER, L_INNER, L_SEP, L_END = layouts[1], layouts[7], layouts[10], layouts[11]  # 版型索引：先列出範本 layouts 再改
X0, W0, Y0, Y1 = 0.37, 9.26, 0.85, 5.15  # 內容區（吋），依範本頁面大小與頁首頁尾調整
orig = list(prs.slides._sldIdLst)


def img(n):
    return glob.glob(f"images/getac/{n}_*.png")[0]


FREE = []
TABLES = []


def slug(title):
    import hashlib
    return hashlib.md5(title.encode("utf-8")).hexdigest()[:8]


def auto_img(title):
    fs = sorted(glob.glob(f"images/getac/auto/a_{slug(title)}_*.png"))
    return fs[-1] if fs else None


def plain(t):
    return t.replace("**", "").replace("`", "")


def emw(c):
    o = ord(c)
    if o > 0x2E7F or 0x2190 <= o <= 0x2BFF: return 1.0
    return 0.3 if c == " " else 0.55


def fit1(t, size, w, floor=11, steps=0):
    """一行寫得完就盡量不換行：不夠寬時最多縮 steps 級（不低於 floor）。"""
    if "\n" in t or lines_of(t, size, w) == 1: return size
    for s_ in range(size - 1, max(size - steps, floor - 1) - 1, -1):
        if s_ >= floor and lines_of(t, s_, w * 0.97) == 1: return s_
    return size


def lines_of(t, size, w):
    t = plain(t)
    out = 0
    for part in t.split("\n"):
        em = w * 72 / size
        out += max(1, math.ceil(sum(emw(c) for c in part) / em * 1.05))
    return out


def rect(s, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = rgb(fill)
    if line: r.line.color.rgb = rgb(line); r.line.width = Pt(0.75)
    else: r.line.fill.background()
    r.shadow.inherit = False
    return r


def runs(para, t, size, color=DK, bold=False):
    for seg in TOK.split(t):
        if not seg: continue
        r = para.add_run()
        b, mono = bold, False
        if seg.startswith("**"): seg = seg[2:-2]; b = True
        elif seg.startswith("`"): seg = seg[1:-1]; mono = True
        r.text = seg
        r.font.name = MONO if mono else FONT
        r.font.size = Pt(size - 1 if mono else size)
        r.font.bold = b
        r.font.color.rgb = rgb(ORG if mono else color)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            e = etree.SubElement(rPr, qn(tag)); e.set("typeface", MONO if mono else FONT)


def text(s, x, y, w, h, t, size=14, bold=False, color=DK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(t if isinstance(t, list) else [t]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = 1.1
        runs(p, line, size, color, bold)
    return tb


def flat(items):
    out = []
    for it in items:
        if isinstance(it, tuple):
            out.append((it[0], 0))
            for sub in it[1]: out.append((sub, 1))
        else:
            out.append((it, 0))
    return out


def bullets_h(items, size, w):
    h = 0
    for t, lv in flat(items):
        aw = w - 0.22 - 0.25 * lv
        sz = fit1(t, size - lv, aw)
        h += lines_of(t, sz, aw) * sz * 1.25 / 72 + 0.07
    return h


def bullets(s, items, x, y, w, size, color=DK):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(bullets_h(items, size, w)))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, (t, lv) in enumerate(flat(items)):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1; p.space_after = Pt(4)
        runs(p, t, fit1(t, size - lv, w - 0.22 - 0.25 * lv), color)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.22 + 0.25 * lv)))); pPr.set("indent", str(-int(Inches(0.2))))
        clr = etree.SubElement(pPr, qn("a:buClr")); c = etree.SubElement(clr, qn("a:srgbClr")); c.set("val", ORG)
        bu = etree.SubElement(pPr, qn("a:buChar")); bu.set("char", "•" if lv == 0 else "–")


def code_h(code, size, w):
    return sum(lines_of(l or " ", size, w - 0.3) for l in code.split("\n")) * size * 1.25 / 72 + 0.25


def code_box(s, code, x, y, w, size):
    h = code_h(code, size, w)
    rect(s, x, y, w, h, LT, BD)
    tb = s.shapes.add_textbox(Inches(x + 0.15), Inches(y + 0.1), Inches(w - 0.3), Inches(h - 0.2))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, l in enumerate(code.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.1
        r = p.add_run(); r.text = l or " "
        r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = rgb(DK)
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:ea", "a:cs"):
            e = etree.SubElement(rPr, qn(tag)); e.set("typeface", FONT)
    return h


def table_geo(rows, widths, size):
    hs = []
    for r in rows:
        hs.append(max(lines_of(c, fit1(c, size, wd - 0.14, 9, 2), wd - 0.14) * fit1(c, size, wd - 0.14, 9, 2) for c, wd in zip(r, widths)) * 1.25 / 72 + 0.1)
    return hs


def table(s, header, rows, widths, x, y, size):
    allrows = [header] + rows
    hs = table_geo(allrows, widths, size)
    gs = s.shapes.add_table(len(allrows), len(widths), Inches(x), Inches(y), Inches(sum(widths)), Inches(sum(hs)))
    tbl = gs.table
    tblPr = gs._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0"); tblPr.set("bandRow", "0")
    for j, wd in enumerate(widths): tbl.columns[j].width = Inches(wd)
    for i, r in enumerate(allrows):
        tbl.rows[i].height = Inches(hs[i])
        for j, c in enumerate(r):
            cell = tbl.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07); cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(DK if i == 0 else (WH if i % 2 else LT))
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.line_spacing = 1.05
            runs(p, c, fit1(c, size, widths[j] - 0.14, 9, 2), WH if i == 0 else DK, i == 0)
    TABLES.append((s, x, y, list(widths), hs))
    return sum(hs)


def table_h(header, rows, widths, size):
    return sum(table_geo([header] + rows, widths, size))


def note_h(note, size, w):
    sz = fit1(note, size, w - 0.5, 10, 2)
    return lines_of(note, sz, w - 0.5) * sz * 1.25 / 72 + 0.25


def note_box(s, note, x, y, w, size, accent=ORG):
    h = note_h(note, size, w)
    rect(s, x, y, w, h, WH, BD); rect(s, x, y, 0.06, h, accent)
    text(s, x + 0.2, y + 0.11, w - 0.4, h - 0.2, note, fit1(note, size, w - 0.5, 10, 2))
    return h


def inner(title):
    s = prs.slides.add_slide(L_INNER)
    s.shapes.title.text = title
    return s


def page(title, items=None, img_name=None, table_=None, code=None, note=None, widths=None, maxsize=16, gap=0.12):
    """items / code / table / note 由上而下排列，統一字級自動縮到放得下。"""
    s = inner(title)
    avail = Y1 - Y0 - 0.05
    wl = W0 - (2.75 if img_name else 0)
    def fits(size):
        parts = []
        if items: parts.append(bullets_h(items, size, wl))
        if code: parts.append(code_h(code, max(size - 3, 9), wl))
        if table_: parts.append(table_h(table_[0], table_[1], widths, size - 1))
        if note: parts.append(note_h(note, size - 1, W0))
        return sum(parts) + gap * (len(parts) - 1) <= avail

    def wraps(size):
        n_ = 0
        for t, lv in flat(items or []):
            aw = wl - 0.22 - 0.25 * lv
            n_ += lines_of(t, fit1(t, size - lv, aw), aw) > 1
        return n_

    best = None   # 優先「換行最少」，同樣少時取字級最大
    for size in range(min(maxsize + 6, 22), 14, -1):
        if fits(size):
            w_ = wraps(size)
            if best is None or w_ < best[1]: best = (size, w_)
    if best: size = best[0]
    else:
        for size in range(13, 8, -1):
            if fits(size): break
    y = Y0
    if items:
        bullets(s, items, X0, y, wl, size); y += bullets_h(items, size, wl) + gap
    if code:
        y += code_box(s, code, X0, y, wl, max(size - 3, 9)) + gap
    if table_:
        y += table(s, table_[0], table_[1], widths, X0, y, size - 1) + gap
    if note:
        y += note_box(s, note, X0, y, W0, size - 1) + gap
    if img_name:
        s.shapes.add_picture(img(img_name), Inches(X0 + wl + 0.1), Inches(Y0 + 0.1), Inches(2.6), Inches(2.6))
    else:
        free = Y1 - y
        FREE.append((title, round(free, 2)))
        auto = auto_img(title)
        if auto and free >= 1.3:
            sz = min(free - 0.1, 2.7)
            s.shapes.add_picture(auto, Inches(X0 + W0 - sz), Inches(Y1 - sz), Inches(sz), Inches(sz))
    return s


def cols_page(title, cols, note=None, accents=None, maxsize=15):
    s = inner(title)
    n = len(cols); gap = 0.15; cw = (W0 - gap * (n - 1)) / n
    avail = Y1 - Y0 - 0.05 - (note_h(note, 12, W0) + 0.15 if note else 0)
    for size in range(min(maxsize + 3, 18), 8, -1):
        hh = max(bullets_h(c[1], size, cw - 0.3) for c in cols) + 0.65
        if hh <= avail: break
    hh = min(avail, max(hh, 1.2))
    for i, (head, items) in enumerate(cols):
        x = X0 + i * (cw + gap)
        rect(s, x, Y0, cw, hh, WH, BD); rect(s, x, Y0, cw, 0.42, ORG if (accents is None or accents[i]) else DK)
        text(s, x + 0.15, Y0, cw - 0.3, 0.42, head, 14, True, WH, anchor=MSO_ANCHOR.MIDDLE)
        bullets(s, items, x + 0.15, Y0 + 0.55, cw - 0.3, size)
    if note: note_box(s, note, X0, Y0 + hh + 0.15, W0, 12)
    return s


def cards_page(title, cards, ncol=2, note=None, maxsize=14):
    s = inner(title)
    nrow = math.ceil(len(cards) / ncol); gap = 0.15; cw = (W0 - gap * (ncol - 1)) / ncol
    avail = Y1 - Y0 - 0.05 - (note_h(note, 12, W0) + 0.15 if note else 0)
    ch = (avail - gap * (nrow - 1)) / nrow
    for size in range(min(maxsize + 3, 18), 8, -1):
        if all(0.5 + (bullets_h(c[1], size, cw - 0.4) if isinstance(c[1], list) else lines_of(c[1], size, cw - 0.4) * size * 1.25 / 72) <= ch for c in cards): break
    for i, (head, body) in enumerate(cards):
        x = X0 + (i % ncol) * (cw + gap); y = Y0 + (i // ncol) * (ch + gap)
        rect(s, x, y, cw, ch, WH, BD); rect(s, x, y, 0.06, ch, ORG if i % 2 == 0 or ncol > 2 else DK)
        text(s, x + 0.2, y + 0.08, cw - 0.35, 0.3, head, 14, True)
        if isinstance(body, list): bullets(s, body, x + 0.2, y + 0.45, cw - 0.4, size)
        else: text(s, x + 0.2, y + 0.45, cw - 0.4, ch - 0.5, body, size)
    if note: note_box(s, note, X0, Y0 + avail + 0.15, W0, 12)
    return s


def separator(title, sub):
    s = prs.slides.add_slide(L_SEP)
    s.shapes.title.text = title
    text(s, 0.35, 3.15, 9.3, 0.4, sub, 14, False, WH, PP_ALIGN.CENTER)
    return s
